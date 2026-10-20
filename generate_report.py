# -*- coding: utf-8 -*-
"""
Script tạo báo cáo chuẩn Đồ án môn học Công nghệ Phần mềm
Trường Đại học Sài Gòn (SGU) - Khoa Công nghệ Thông tin
Đề tài: Hệ thống Thuyết minh Tự động Đa ngôn ngữ (Smart Multilingual Audio Guide)
"""

import os
import sys
import json
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from PIL import Image

# ----------------- HẰNG SỐ MÀU SẮC & PHÔNG CHỮ -----------------
COLOR_PRIMARY_HEX = "003366"      # Xanh Navy ĐH Sài Gòn
COLOR_SECONDARY_HEX = "1A5276"    # Xanh phụ
COLOR_GOLD_HEX = "D4AC0D"         # Vàng điểm xuyết
COLOR_GRAY_HEX = "555555"         # Xám chú thích
COLOR_BORDER_HEX = "B0C4DE"       # Xám xanh viền bảng
COLOR_BG_HEADER_HEX = "003366"    # Nền tiêu đề bảng
COLOR_ZEBRA_HEX = "F4F6F9"        # Nền sọc bảng

COLOR_PRIMARY = RGBColor(0, 51, 102)
COLOR_SECONDARY = RGBColor(26, 82, 118)
COLOR_TEXT = RGBColor(33, 37, 41)
COLOR_MUTED = RGBColor(108, 117, 125)
COLOR_GOLD = RGBColor(212, 172, 13)

# ----------------- HÀM TIỆN ÍCH ĐỊNH DẠNG WORD -----------------

def set_cell_background(cell, hex_color):
    """Đặt màu nền cho cell của bảng"""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Đặt padding cho ô trong bảng (đơn vị dxa: 20 dxa = 1 pt)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_borders(cell, top="single", bottom="single", left="none", right="none",
                     sz="4", color="B0C4DE"):
    """Đặt viền thanh lịch cho từng cell"""
    tcPr = cell._tc.get_or_add_tcPr()
    xml = f'<w:tcBorders {nsdecls("w")}>'
    xml += f'<w:top w:val="{top}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
    xml += f'<w:bottom w:val="{bottom}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
    xml += f'<w:left w:val="{left}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
    xml += f'<w:right w:val="{right}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
    xml += '</w:tcBorders>'
    tcPr.append(parse_xml(xml))

def add_page_number_to_run(run):
    """Thêm mã trường số trang động PAGE vào Run"""
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)

def add_numpages_to_run(run):
    """Thêm mã trường tổng số trang NUMPAGES vào Run"""
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'NUMPAGES'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)

def style_paragraph(p, space_before=0, space_after=6, line_spacing=1.3, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    """Định dạng paragraph chuẩn chỉ"""
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.alignment = align

def add_heading_1(doc, text):
    """Thêm Heading 1 chuẩn phong cách Báo cáo SGU"""
    h = doc.add_paragraph()
    style_paragraph(h, space_before=16, space_after=8, align=WD_ALIGN_PARAGRAPH.LEFT)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    return h

def add_heading_2(doc, text):
    """Thêm Heading 2"""
    h = doc.add_paragraph()
    style_paragraph(h, space_before=12, space_after=6, align=WD_ALIGN_PARAGRAPH.LEFT)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = COLOR_SECONDARY
    return h

def add_heading_3(doc, text):
    """Thêm Heading 3"""
    h = doc.add_paragraph()
    style_paragraph(h, space_before=8, space_after=4, align=WD_ALIGN_PARAGRAPH.LEFT)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = COLOR_TEXT
    return h

def add_body_p(doc, text, bold_prefix=None, italic=False):
    """Thêm đoạn văn bản nội dung thông thường"""
    p = doc.add_paragraph()
    style_paragraph(p, space_before=0, space_after=6, line_spacing=1.3, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(13)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_TEXT
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.italic = italic
    r.font.color.rgb = COLOR_TEXT
    return p

def add_bullet_p(doc, text, bold_prefix=None, level=0):
    """Thêm dòng gạch đầu dòng chuẩn"""
    p = doc.add_paragraph(style='List Bullet')
    style_paragraph(p, space_before=0, space_after=4, line_spacing=1.25, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(13)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_TEXT
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.color.rgb = COLOR_TEXT
    return p

def add_caption(doc, text, kind="Hình"):
    """Thêm ghi chú hình ảnh hoặc bảng biểu theo chuẩn học thuật"""
    p = doc.add_paragraph()
    style_paragraph(p, space_before=4, space_after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(f"{kind}: {text}")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r.font.italic = True
    r.font.bold = True
    r.font.color.rgb = COLOR_GRAY_HEX and COLOR_PRIMARY
    return p

def insert_image_safely(doc, img_path, caption_text, max_width_in=6.0, max_height_in=6.5):
    """Chèn hình ảnh an toàn, tự động tính tỷ lệ co giãn"""
    if not os.path.exists(img_path):
        print(f"Warning: Image not found: {img_path}")
        return
    try:
        with Image.open(img_path) as im:
            w, h = im.size
        # Tính tỷ lệ inches
        aspect = h / w
        target_w = max_width_in
        target_h = target_w * aspect
        if target_h > max_height_in:
            target_h = max_height_in
            target_w = target_h / aspect

        p = doc.add_paragraph()
        style_paragraph(p, space_before=8, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
        p.paragraph_format.keep_with_next = True
        run = p.add_run()
        run.add_picture(img_path, width=Inches(target_w), height=Inches(target_h))
        add_caption(doc, caption_text, kind="Hình")
    except Exception as e:
        print(f"Error inserting image {img_path}: {e}")

print("Helper definitions loaded successfully.")
