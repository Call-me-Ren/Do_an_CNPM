import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def sync_dactausecase():
    path = 'docs/requirements/assets/dactausecase.docx'
    doc = docx.Document(path)
    
    # 1. Update Table 1 (USC01)
    t1 = doc.tables[1]
    for row in t1.rows:
        first_txt = row.cells[0].text.strip()
        if 'Luồng ngoại lệ' in first_txt:
            row.cells[1].text = 'E1 - Sai mật khẩu: Hệ thống thông báo "Tên đăng nhập hoặc mật khẩu không chính xác" và cho phép nhập lại (tự động khóa tạm thời 5 phút nếu nhập sai quá 5 lần liên tiếp).'
        elif 'Điều kiện kết thúc' in first_txt:
            row.cells[1].text = 'Thành công: Phiên làm việc (Session/Token) được khởi tạo thành công với quyền hạn tương ứng (Admin truy cập Dashboard Quản trị hệ thống; khách tham quan sử dụng trực tiếp không cần tài khoản).\nThất bại: Hệ thống từ chối quyền truy cập, hiển thị thông báo lỗi (E1) và không tạo phiên làm việc.'

    # 2. Update Table 5 (USC05)
    t5 = doc.tables[5]
    for row in t5.rows:
        first_txt = row.cells[0].text.strip()
        if 'Dòng sự kiện chính' in first_txt:
            row.cells[1].text = '1. Người dùng xem chi tiết hiện vật (không yêu cầu đăng nhập).\n2. Hệ thống hiển thị nút "Lưu Yêu Thích" (biểu tượng trái tim).\n3. Người dùng bấm nút "Lưu Yêu Thích".\n4. Hệ thống kiểm tra: nếu chưa lưu thì thêm hiện vật vào danh sách yêu thích trên thiết bị (LocalStorage); nếu đã lưu thì hiển thị tùy chọn hủy lưu (A1).\n5. Hệ thống cập nhật biểu tượng trạng thái và thông báo "Đã lưu vào danh sách yêu thích".'
        elif 'Điều kiện kết thúc' in first_txt:
            row.cells[1].text = 'Hiện vật được lưu vào danh sách yêu thích trực tiếp trên thiết bị của người dùng (hỗ trợ lưu bền vững qua LocalStorage).'

    # 3. Update Table 10 (USC10)
    t10 = doc.tables[10]
    t10.rows[0].cells[0].text = 'USC10: Quản Lý Phiên Tham Quan & Mã QR (QuanLyPhienVaMaQR)'
    if len(t10.rows[0].cells) > 1:
        t10.rows[0].cells[1].text = 'USC10: Quản Lý Phiên Tham Quan & Mã QR (QuanLyPhienVaMaQR)'
    for row in t10.rows:
        first_txt = row.cells[0].text.strip()
        if 'Mô tả tóm tắt' in first_txt:
            row.cells[1].text = 'Cho phép Admin tạo mới, chỉnh sửa, xóa và theo dõi các phiên tham quan (ca mở cửa, đoàn tham quan chuyên biệt), đồng thời cấu hình và xuất mã QR tương ứng cho cổng đón tiếp và từng hiện vật. Mỗi phiên có mã riêng, khung thời gian, ghi chú và số liệu số khách, số lượt nghe, giúp hệ thống thống kê lượt nghe chính xác theo từng ca và từng đoàn (USC12).'

    # 4. Update Table 14 (USC14)
    t14 = doc.tables[14]
    t14.rows[0].cells[0].text = 'USC14: Định Vị Vị Trí & Tự Động Thuyết Minh (DinhViViTriVaTuDongThuyetMinh)'
    if len(t14.rows[0].cells) > 1:
        t14.rows[0].cells[1].text = 'USC14: Định Vị Vị Trí & Tự Động Thuyết Minh (DinhViViTriVaTuDongThuyetMinh)'
    for row in t14.rows:
        first_txt = row.cells[0].text.strip()
        if 'Mô tả tóm tắt' in first_txt:
            row.cells[1].text = 'Xác định và cập nhật liên tục vị trí của người dùng trong khu trưng bày thông qua cơ chế lai: quét mã QR tại cổng hoặc hiện vật làm mốc neo xuất phát (anchorArtifactId) và tự động kích hoạt GPS theo dõi độ chính xác cao. Hệ thống đối chiếu tọa độ với bán kính kích hoạt của các hiện vật (mặc định 12m), hiển thị pin "Bạn đang ở đây" trên sơ đồ tầng, và tự động đề xuất/kích hoạt thuyết minh tương ứng khi đến gần mà không cần quét lại.'

    # 5. Update Paragraph Headings
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt == '3.10 USC10 — Quản lý phiên':
            p.text = '3.10 USC10 — Quản Lý Phiên Tham Quan & Mã QR'
        elif txt == '3.14 USC14 — Định vị vị trí':
            p.text = '3.14 USC14 — Định Vị Vị Trí & Tự Động Thuyết Minh'
        elif '%3CmxGraphModel%3E' in txt:
            p.text = '' # Clear garbage xml string

    # 6. Add USC15, Component Diagram, Deployment Diagram
    # Heading USC15
    p_u15 = doc.add_paragraph('3.15 USC15 — Lưu Trữ & Phân Phối Tệp Âm Thanh', style='Heading 2')
    p_u15_diag = doc.add_paragraph('Sơ đồ Use Case — USC15')
    p_u15_diag.runs[0].italic = True
    
    # Table USC15
    t15 = doc.add_table(rows=10, cols=2)
    t15.style = 'Normal Table'
    t15.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    data_u15 = [
        ('USC15: Lưu Trữ & Phân Phối Tệp Âm Thanh (LuuTruVaPhanPhoiTepAmThanh)', 'USC15: Lưu Trữ & Phân Phối Tệp Âm Thanh (LuuTruVaPhanPhoiTepAmThanh)'),
        ('Tác nhân (Actor):', 'Admin, Thiết Bị Lưu Âm Thanh (Web CDN / Static Audio Server)'),
        ('Mức độ hoàn thiện:', 'Focused'),
        ('Mô tả tóm tắt:', 'Lưu trữ và phân phối tệp âm thanh thuyết minh tĩnh (kèm thông tin hiện vật, ngôn ngữ) lên Thiết Bị Lưu Âm Thanh / Web CDN để phục vụ việc phát luồng qua giao thức HTTP/HTTPS. Use case này phối hợp cùng {Quản Lý Tệp Âm Thanh} («include») để nhận tệp và thông tin do Admin cung cấp.'),
        ('Điều kiện tiên quyết:', '- Thiết Bị Lưu Âm Thanh / Web CDN sẵn sàng tiếp nhận và phục vụ tệp.\n- Tệp âm thanh và thông tin kèm theo đã được Admin cung cấp qua {Quản Lý Tệp Âm Thanh}.'),
        ('Dòng sự kiện chính:', '1. Hệ thống tiếp nhận tệp âm thanh thuyết minh và thông tin kèm theo (hiện vật, ngôn ngữ) từ use case {Quản Lý Tệp Âm Thanh}.\n2. Hệ thống kiểm tra tệp: đúng định dạng âm thanh (.mp3) và dung lượng trong giới hạn cho phép (E1).\n3. Hệ thống kiểm tra tệp đã tồn tại chưa (A1), tạo ID tệp và gửi yêu cầu lưu kèm dữ liệu đến kho lưu trữ.\n4. Thiết Bị Lưu Âm Thanh / Kho lưu trữ kiểm tra hạn mức dung lượng trống (E2).\n5. Thiết bị ghi tệp lên bộ nhớ và kiểm tra tính toàn vẹn (E3).\n6. Thiết bị trả kết quả thành công kèm đường dẫn URL tệp audio.\n7. Hệ thống cập nhật thông tin đường dẫn tệp vào CSDL/Store và thông báo "Lưu tệp âm thanh thành công". Use case kết thúc.'),
        ('Luồng thay thế (Alternative):', 'A1 - Ghi đè/thay thế tệp đã có: Tại bước 3, nếu đã tồn tại tệp cùng hiện vật và ngôn ngữ, hệ thống cho phép ghi đè tệp mới; tệp cũ chỉ bị thay thế sau khi tệp mới lưu thành công.'),
        ('Luồng ngoại lệ (Exception):', 'E1 - Tệp không hợp lệ: Sai định dạng âm thanh hoặc vượt dung lượng cho phép. Hệ thống báo lỗi và hủy thao tác.\nE2 - Hết dung lượng lưu trữ: Hệ thống thông báo lỗi không đủ dung lượng.\nE3 - Ghi tệp thất bại hoặc mất kết nối: Hệ thống tự động thử lại, nếu không thành công thì hủy phần tệp ghi dở và báo lỗi.'),
        ('Điều kiện kết thúc:', 'Thành công: Tệp âm thanh được lưu đầy đủ trên kho lưu trữ và đường dẫn tệp được cập nhật vào CSDL/Store.\nThất bại: Không lưu được tệp; hệ thống không giữ lại tệp ghi dở.'),
        ('Sự kiện kích hoạt (Trigger):', 'Admin thêm mới hoặc cập nhật tệp âm thanh thuyết minh cho hiện vật.')
    ]
    
    for idx, (col0, col1) in enumerate(data_u15):
        row = t15.rows[idx]
        row.cells[0].text = col0
        row.cells[1].text = col1
        # Style row 0 header
        if idx == 0:
            row.cells[0].paragraphs[0].runs[0].bold = True
            row.cells[1].paragraphs[0].runs[0].bold = True

    p_u15_act = doc.add_paragraph('Sơ đồ hoạt động (Activity Diagram) — USC15')
    p_u15_act.runs[0].italic = True
    doc.add_paragraph('Tiến trình: Nhận tệp âm thanh & metadata -> Kiểm tra định dạng audio -> Kiểm tra hạn mức dung lượng -> Ghi tệp toàn vẹn -> Cập nhật URL vào CSDL / Data Store -> Thông báo thành công.')

    # Section 4: Component Diagram
    p_sec4 = doc.add_paragraph('4. Sơ Đồ Thành Phần (Component Diagram)', style='Heading 1')
    p_sec4_desc = doc.add_paragraph(
        'Sơ đồ thành phần mô tả cấu trúc hệ thống theo kiến trúc SPA Front-end hiện đại và các dịch vụ ngoại vi:\n'
        '1. Phân hệ Khách Tham Quan (Visitor App): Gồm VisitorLayout (Router QR & GPS), TrackingPanel (theo dõi vị trí & đề xuất chuyển/nghe lại), '
        'AudioPlayer & Subtitles (phát audio & phụ đề karaoke từng giây), FloorMapView (bản đồ tương tác đa tầng & halo vị trí "Bạn đang ở đây"), '
        'ExploreView (tìm kiếm theo tên, mã số ART-xxx, mô tả), ArtifactDetailView (nội dung chi tiết hiện vật).\n'
        '2. Phân hệ Quản Trị (Admin App): Gồm AdminLoginView (bảo mật khóa tạm thời 5 phút khi nhập sai quá 5 lần liên tiếp), '
        'AdminDashboard (thống kê lượt nghe theo phiên và hiện vật), ArtifactsManageView (CRUD hiện vật, cấu hình tọa độ GPS vĩ độ/kinh độ & bán kính kích hoạt triggerRadius), '
        'SessionsManageView (quản lý phiên tham quan & xuất mã QR đón tiếp tại cổng).\n'
        '3. Khối Core Modules & Shared Services: geofence.ts (tính khoảng cách Haversine & vùng kích hoạt), qrHelper.ts (giải mã QR & deep link), '
        'loginLock.ts (động cơ bảo mật lockout), các Zustand stores (geofenceStore, audioStore, favoritesStore, languageStore, sessionStore) '
        'và api.ts (mock data & lưu trữ LocalStorage bền vững).\n'
        '4. Hạ tầng Lưu trữ & Dịch vụ: GitHub Pages (lưu trữ mã nguồn tĩnh HTML/JS/CSS, âm thanh MP3 và hình ảnh sơ đồ) '
        'kết hợp dịch vụ ntfy.sh phục vụ giám sát và telemetry thời gian thực.'
    )

    # Section 5: Deployment Diagram
    p_sec5 = doc.add_paragraph('5. Sơ Đồ Triển Khai (Deployment Diagram)', style='Heading 1')
    p_sec5_desc = doc.add_paragraph(
        'Sơ đồ triển khai xác định môi trường thực thi và hạ tầng mạng của hệ thống:\n'
        '1. Nút Thiết bị Người dùng (Client Device - Mobile / Desktop): Chạy trên trình duyệt Web hiện đại (Chrome, Safari, Edge, Firefox), '
        'tận dụng trực tiếp các Web API chuẩn mực: HTML5 Web Audio API (phát luồng âm thanh), HTML5 Geolocation API (định vị GPS độ chính xác cao), '
        'HTML5 MediaDevices / Camera API (quét mã QR trực tiếp qua camera), và Browser LocalStorage (lưu hiện vật yêu thích, token và cài đặt người dùng mà không cần cài đặt ứng dụng).\n'
        '2. Nút Hạ tầng GitHub Pages (Nginx CDN HTTPS :443): Đóng gói bản dựng Vite SPA, tệp chỉ mục index.html, các tài sản âm thanh public/audio/*.mp3 '
        'và hình ảnh sơ đồ public/images/*.png, phân phối toàn cầu với độ trễ thấp và bảo mật SSL/TLS.\n'
        '3. Nút Hạ tầng Giám Sát / Analytics (ntfy.sh HTTPS Gateway): Tiếp nhận các sự kiện thống kê, lượt nghe và phiên tham quan qua giao thức HTTPS POST.'
    )

    doc.save(path)
    print('dactausecase.docx synchronized successfully!')

if __name__ == '__main__':
    sync_dactausecase()
